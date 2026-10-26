from io import BytesIO
from pathlib import Path
import csv, io
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED = {'.pdf','.csv','.xlsx','.xlsm','.docx','.txt','.md','.html','.htm'}

def _pdf_text(data: bytes) -> str:
    return '\n'.join((p.extract_text() or '') for p in PdfReader(BytesIO(data)).pages)

def _ocr_pdf(data: bytes) -> str:
    import fitz, pytesseract
    from PIL import Image
    doc = fitz.open(stream=data, filetype='pdf'); pages=[]
    for page in doc:
        pix = page.get_pixmap(matrix=fitz.Matrix(2,2), alpha=False)
        img = Image.frombytes('RGB', [pix.width,pix.height], pix.samples)
        pages.append(pytesseract.image_to_string(img, lang='eng+kan+hin+tel'))
    return '\n'.join(pages)

def extract(name: str, data: bytes) -> tuple[str,bool]:
    ext=Path(name).suffix.lower()
    if ext not in SUPPORTED: raise ValueError('Unsupported file format.')
    if ext=='.pdf':
        text=_pdf_text(data).strip()
        if len(text) >= 80: return text,False
        return _ocr_pdf(data).strip(),True
    if ext=='.docx': return '\n'.join(p.text for p in DocxDocument(BytesIO(data)).paragraphs),False
    if ext in {'.xlsx','.xlsm'}:
        wb=load_workbook(BytesIO(data),read_only=True,data_only=True)
        return '\n'.join(f'[{ws.title}] '+' | '.join('' if v is None else str(v) for v in row) for ws in wb.worksheets for row in ws.iter_rows(values_only=True)),False
    text=data.decode('utf-8-sig',errors='replace')
    if ext=='.csv': return '\n'.join(' | '.join(row) for row in csv.reader(io.StringIO(text))),False
    if ext in {'.html','.htm'}: return BeautifulSoup(text,'html.parser').get_text('\n',strip=True),False
    return text,False
